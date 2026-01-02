#!/usr/bin/env python3
"""
Python 3.11 equivalent pipeline for the n8n workflow: AnalisisDataPerusahaanTbk.

Source workflow JSON: AnalisisDataPerusahaanTbk.json
- Webhook Trigger (POST /analisa-emiten)
- Weekly schedule (Monday 06:00 Asia/Jakarta)
- Data sources: Google Sheets (stock prices), Google News RSS, Annual Reports (PDF URLs), optional Bond API
- Analytics: 5Y stock metrics, financial ratio extraction (heuristic), news->SWOT classification
- Report: OpenAI structured output -> PDF/DOCX -> optional Google Drive upload + email + Sheets log

Notes:
- This is an engineering translation of the workflow logic; it is not a 1:1 node runtime.
- Several integrations (Google APIs, bond API, SMTP, OpenAI) require credentials via environment variables.
- Financial ratio extraction from PDFs is heuristic and should be validated against audited statements.
"""

from __future__ import annotations

import argparse
import dataclasses
import datetime as dt
import json
import logging
import math
import os
import re
import statistics
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import requests
import feedparser
from dateutil import tz
from pydantic import BaseModel, Field, ValidationError

# Optional: FastAPI + scheduler (only needed when serving)
try:
    from fastapi import FastAPI, HTTPException
    from fastapi.responses import JSONResponse
    from apscheduler.schedulers.background import BackgroundScheduler
    from apscheduler.triggers.cron import CronTrigger
except Exception:  # pragma: no cover
    FastAPI = None  # type: ignore


LOGGER = logging.getLogger("analisis_emiten")
TZ_JAKARTA = tz.gettz("Asia/Jakarta")


# -----------------------------
# Models (inputs + OpenAI output)
# -----------------------------

class AnalysisRequest(BaseModel):
    company_name: str = Field(default="<__PLACEHOLDER_VALUE__Nama Perusahaan__>")
    stock_ticker: str = Field(default="<__PLACEHOLDER_VALUE__Ticker Saham (mis: IDX:BBCA)__>")
    period_years: int = Field(default=5, ge=1, le=20)
    benchmark_ticker: str = Field(default="IHSG")
    bond_identifiers: List[str] = Field(default_factory=list)
    annual_report_urls: List[str] = Field(default_factory=list)
    news_query: Optional[str] = None
    language: str = Field(default="id")
    output_format: List[str] = Field(default_factory=lambda: ["PDF", "JSON"])
    google_sheet_id: str = Field(default="<__PLACEHOLDER_VALUE__Google Sheet ID untuk data saham__>")
    drive_folder_id: str = Field(default="<__PLACEHOLDER_VALUE__Google Drive Folder ID__>")
    recipients_email: List[str] = Field(default_factory=list)

    def normalized(self) -> "AnalysisRequest":
        if not self.news_query:
            self.news_query = self.company_name
        return self


class StockPoint(BaseModel):
    date: dt.date
    close: float
    volume: float = 0.0


class NewsItem(BaseModel):
    title: str
    url: str
    source: str = "Unknown"
    date: dt.datetime
    snippet: str = ""


class ClassifiedNewsItem(BaseModel):
    title: str
    description: str
    link: str
    pubDate: str
    swotCategory: str
    theme: str
    sentimentScore: float
    materialityScore: float


class ReportStockPerformance(BaseModel):
    narrative: str
    metrics: Dict[str, Any]


class ReportBondView(BaseModel):
    narrative: str
    limitations: str


class ReportFinancialAnalysis(BaseModel):
    highlights: str
    tables: List[Dict[str, Any]]


class ReportSWOT(BaseModel):
    S: List[str]
    W: List[str]
    O: List[str]
    T: List[str]


class ReportAppendix(BaseModel):
    sources: List[str]
    assumptions: List[str]
    methodology: str


class Big4Report(BaseModel):
    executive_summary: str
    company_overview: str
    industry: str
    stock_performance: ReportStockPerformance
    bond_view: ReportBondView
    financial_analysis: ReportFinancialAnalysis
    swot: ReportSWOT
    key_risks: List[str]
    governance_esg_regulatory_notes: str
    appendix: ReportAppendix


# -----------------------------
# Utility
# -----------------------------

def now_jakarta() -> dt.datetime:
    return dt.datetime.now(tz=TZ_JAKARTA)


def run_id(stock_ticker: str) -> str:
    return f"{now_jakarta().strftime('%Y%m%d-%H%M%S')}-{stock_ticker}"


def safe_float(x: Any, default: float = 0.0) -> float:
    try:
        if x is None:
            return default
        if isinstance(x, (int, float)):
            return float(x)
        s = str(x).strip()
        if s == "":
            return default
        # strip thousands separators
        s = s.replace(",", "")
        return float(s)
    except Exception:
        return default


def parse_date_maybe(value: Any) -> Optional[dt.date]:
    """Accepts Date objects, strings, or Excel serial numbers."""
    if value is None:
        return None
    if isinstance(value, dt.date) and not isinstance(value, dt.datetime):
        return value
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, (int, float)):
        # Excel serial date: 25569 -> 1970-01-01
        try:
            epoch = dt.date(1970, 1, 1)
            return epoch + dt.timedelta(days=float(value) - 25569)
        except Exception:
            return None
    if isinstance(value, str):
        s = value.strip()
        if s.lower() in {"date", "tanggal"}:
            return None
        # Try ISO first
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y"):
            try:
                return dt.datetime.strptime(s, fmt).date()
            except Exception:
                pass
        # Fallback to dateutil parser
        try:
            from dateutil.parser import parse
            return parse(s).date()
        except Exception:
            return None
    return None


def iqr_outlier_filter(points: List[StockPoint], field: str = "close") -> List[StockPoint]:
    values = sorted([getattr(p, field) for p in points if getattr(p, field) is not None and not math.isnan(getattr(p, field))])
    if len(values) < 4:
        return points
    q1 = values[math.floor(len(values) * 0.25)]
    q3 = values[math.floor(len(values) * 0.75)]
    iqr = q3 - q1
    lo = q1 - 1.5 * iqr
    hi = q3 + 1.5 * iqr
    return [p for p in points if lo <= getattr(p, field) <= hi]


# -----------------------------
# Google Sheets (Stock Data)
# -----------------------------

def _get_google_creds():
    """
    Uses Application Default Credentials (recommended):
      export GOOGLE_APPLICATION_CREDENTIALS=/path/service_account.json
    """
    try:
        from google.auth import default
        creds, _ = default(scopes=[
            "https://www.googleapis.com/auth/spreadsheets",
            "https://www.googleapis.com/auth/drive",
        ])
        return creds
    except Exception as e:
        raise RuntimeError(
            "Google credentials not configured. Set GOOGLE_APPLICATION_CREDENTIALS "
            "or install/authorize appropriate credentials."
        ) from e


def read_stock_data_google_sheets(sheet_id: str, sheet_name: str = "Stock Data") -> List[Dict[str, Any]]:
    """
    Expects columns that include at least:
      - date (Date)
      - close (Close/close/price/Price)
      - volume (optional)
    """
    creds = _get_google_creds()
    from googleapiclient.discovery import build
    service = build("sheets", "v4", credentials=creds, cache_discovery=False)

    # Read a broad range; adjust as needed
    rng = f"{sheet_name}!A:Z"
    res = service.spreadsheets().values().get(spreadsheetId=sheet_id, range=rng).execute()
    values = res.get("values", [])
    if not values:
        return []

    headers = [h.strip() for h in values[0]]
    rows = values[1:]
    data = []
    for r in rows:
        row = {headers[i]: (r[i] if i < len(r) else "") for i in range(len(headers))}
        data.append(row)
    return data


# -----------------------------
# News (Google News RSS)
# -----------------------------

def fetch_google_news_rss(query: str, hl: str = "id", gl: str = "ID", ceid: str = "ID:id", limit: int = 50) -> List[NewsItem]:
    url = f"https://news.google.com/rss/search?q={requests.utils.quote(query)}&hl={hl}&gl={gl}&ceid={ceid}"
    feed = feedparser.parse(url)
    items: List[NewsItem] = []
    seen = set()

    for entry in feed.entries[: max(limit, 1)]:
        link = getattr(entry, "link", "") or ""
        if not link or link in seen:
            continue
        seen.add(link)

        title = (getattr(entry, "title", "") or "").strip()
        snippet = (getattr(entry, "summary", "") or getattr(entry, "description", "") or "")
        snippet = re.sub(r"<[^>]*>", "", snippet).replace("&amp;", "&").replace("&nbsp;", " ").strip()

        source = getattr(entry, "source", None)
        source_name = getattr(source, "title", None) if source else None
        source_str = source_name or getattr(entry, "author", None) or "Unknown"

        dt_parsed = None
        if getattr(entry, "published_parsed", None):
            dt_parsed = dt.datetime(*entry.published_parsed[:6], tzinfo=dt.timezone.utc)
        elif getattr(entry, "updated_parsed", None):
            dt_parsed = dt.datetime(*entry.updated_parsed[:6], tzinfo=dt.timezone.utc)
        else:
            dt_parsed = dt.datetime.now(dt.timezone.utc)

        items.append(NewsItem(
            title=html_unescape(title),
            url=link,
            source=html_unescape(str(source_str)),
            date=dt_parsed.astimezone(TZ_JAKARTA),
            snippet=html_unescape(snippet),
        ))

    # newest first
    items.sort(key=lambda x: x.date, reverse=True)
    return items


def html_unescape(text: str) -> str:
    return (
        text.replace("&amp;", "&")
            .replace("&lt;", "<")
            .replace("&gt;", ">")
            .replace("&quot;", '"')
            .replace("&#39;", "'")
            .replace("&nbsp;", " ")
            .strip()
    )


# -----------------------------
# Annual Reports (PDF) -> text
# -----------------------------

def download_files(urls: List[str], out_dir: Path, max_files: int = 5, timeout: int = 60) -> List[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths: List[Path] = []
    for i, u in enumerate(urls[:max_files], start=1):
        if not u:
            continue
        fn = out_dir / f"annual_report_{i}.pdf"
        LOGGER.info("Downloading annual report %s -> %s", u, fn)
        r = requests.get(u, timeout=timeout)
        r.raise_for_status()
        fn.write_bytes(r.content)
        paths.append(fn)
    return paths


def extract_text_from_pdfs(pdf_paths: List[Path]) -> List[str]:
    """
    Uses pypdf (pip install pypdf).
    """
    texts: List[str] = []
    try:
        from pypdf import PdfReader
    except Exception as e:
        raise RuntimeError("Missing dependency: pypdf") from e

    for p in pdf_paths:
        reader = PdfReader(str(p))
        parts = []
        for page in reader.pages:
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                continue
        texts.append("\n".join(parts))
    return texts


# -----------------------------
# Parse & Clean Stock Data
# -----------------------------

def parse_and_clean_stock_data(rows: List[Dict[str, Any]]) -> Tuple[List[StockPoint], Dict[str, Any]]:
    cleaned: List[StockPoint] = []

    for row in rows:
        # Try multiple common header variants
        date_val = row.get("date") or row.get("Date") or row.get("tanggal") or row.get("Tanggal")
        d = parse_date_maybe(date_val)
        if not d:
            continue

        close_val = row.get("close") or row.get("Close") or row.get("price") or row.get("Price") or row.get("CLOSE")
        close = safe_float(close_val, default=float("nan"))
        if math.isnan(close) or close <= 0:
            continue

        vol_val = row.get("volume") or row.get("Volume") or 0
        vol = safe_float(vol_val, default=0.0)

        cleaned.append(StockPoint(date=d, close=close, volume=vol))

    cleaned.sort(key=lambda x: x.date)

    filtered = iqr_outlier_filter(cleaned, "close")
    close_prices = [p.close for p in filtered]
    volumes = [p.volume for p in filtered]

    stats: Dict[str, Any] = {}
    if filtered:
        stats = {
            "count": len(filtered),
            "priceStats": {
                "min": min(close_prices),
                "max": max(close_prices),
                "mean": statistics.fmean(close_prices),
                "median": statistics.median(close_prices),
            },
            "volumeStats": {
                "min": min(volumes) if volumes else 0,
                "max": max(volumes) if volumes else 0,
                "mean": statistics.fmean(volumes) if volumes else 0,
                "total": sum(volumes) if volumes else 0,
            },
            "dateRange": {"start": filtered[0].date.isoformat(), "end": filtered[-1].date.isoformat()},
            "outliersRemoved": len(cleaned) - len(filtered),
        }

    return filtered, stats


# -----------------------------
# Stock Analytics (5Y)
# -----------------------------

def _returns(prices: List[float]) -> List[float]:
    r: List[float] = []
    for i in range(1, len(prices)):
        p0 = prices[i - 1]
        p1 = prices[i]
        if p0:
            r.append((p1 - p0) / p0)
    return r


def _cagr(start: float, end: float, years: float) -> float:
    if start <= 0 or end <= 0 or years <= 0:
        return 0.0
    return (end / start) ** (1 / years) - 1


def _volatility(daily_returns: List[float]) -> float:
    if not daily_returns:
        return 0.0
    mean = statistics.fmean(daily_returns)
    var = statistics.fmean([(x - mean) ** 2 for x in daily_returns])
    daily = math.sqrt(var)
    return daily * math.sqrt(252)  # annualized


def _max_drawdown(prices: List[float], dates: List[dt.date]) -> Dict[str, Any]:
    peak = prices[0]
    peak_date = dates[0]
    mdd = 0.0
    trough_date = dates[0]
    for p, d in zip(prices, dates):
        if p > peak:
            peak = p
            peak_date = d
        dd = (peak - p) / peak if peak else 0.0
        if dd > mdd:
            mdd = dd
            trough_date = d
    return {"maxDrawdown": mdd, "peakDate": peak_date.isoformat(), "troughDate": trough_date.isoformat()}


def _moving_average(prices: List[float], period: int) -> List[float]:
    if period <= 0 or len(prices) < period:
        return []
    out = []
    window_sum = sum(prices[:period])
    out.append(window_sum / period)
    for i in range(period, len(prices)):
        window_sum += prices[i] - prices[i - period]
        out.append(window_sum / period)
    return out


def _rolling_return(prices: List[float], period: int) -> List[float]:
    if period <= 0 or len(prices) <= period:
        return []
    out = []
    for i in range(period, len(prices)):
        p0 = prices[i - period]
        p1 = prices[i]
        out.append((p1 - p0) / p0 if p0 else 0.0)
    return out


def _correlation(a: List[float], b: List[float]) -> float:
    n = min(len(a), len(b))
    if n == 0:
        return 0.0
    a = a[:n]
    b = b[:n]
    ma = statistics.fmean(a)
    mb = statistics.fmean(b)
    num = sum((x - ma) * (y - mb) for x, y in zip(a, b))
    den = math.sqrt(sum((x - ma) ** 2 for x in a) * sum((y - mb) ** 2 for y in b))
    return (num / den) if den else 0.0


def calculate_stock_analytics(
    stock: List[StockPoint],
    period_years: int,
    benchmark: Optional[List[StockPoint]] = None,
    risk_free_rate: float = 0.06,
) -> Dict[str, Any]:
    if not stock:
        return {"error": "No stock data"}

    prices = [p.close for p in stock]
    dates = [p.date for p in stock]
    start_price, end_price = prices[0], prices[-1]

    years = float(period_years)
    cagr = _cagr(start_price, end_price, years)
    total_return = (end_price - start_price) / start_price if start_price else 0.0

    rets = _returns(prices)
    vol = _volatility(rets)
    mdd = _max_drawdown(prices, dates)

    rolling_1y = _rolling_return(prices, 252)
    rolling_3y = _rolling_return(prices, 252 * 3)
    ma50 = _moving_average(prices, 50)
    ma200 = _moving_average(prices, 200)

    current_ma50 = ma50[-1] if ma50 else None
    current_ma200 = ma200[-1] if ma200 else None

    # Benchmark
    bench_prices: List[float] = []
    bench_rets: List[float] = []
    bench_cagr = 0.0
    corr = 0.0
    beta = 1.0
    outperformance = None

    if benchmark:
        bench_prices = [p.close for p in benchmark]
        bench_rets = _returns(bench_prices)
        bench_cagr = _cagr(bench_prices[0], bench_prices[-1], years) if bench_prices else 0.0
        corr = _correlation(rets, bench_rets)
        bench_vol = _volatility(bench_rets)
        beta = (corr * vol) / bench_vol if bench_vol else 1.0
        outperformance = cagr - bench_cagr

    sharpe = (cagr - risk_free_rate) / vol if vol else 0.0

    return {
        "analysisDate": now_jakarta().isoformat(),
        "periodYears": period_years,
        "priceMetrics": {
            "startPrice": start_price,
            "endPrice": end_price,
            "startDate": dates[0].isoformat(),
            "endDate": dates[-1].isoformat(),
        },
        "returnMetrics": {
            "cagr": cagr,
            "cagrPercent": f"{cagr * 100:.2f}%",
            "totalReturn": total_return,
            "totalReturnPercent": f"{total_return * 100:.2f}%",
            "rolling1YAvg": statistics.fmean(rolling_1y) if rolling_1y else 0.0,
            "rolling3YAvg": statistics.fmean(rolling_3y) if rolling_3y else 0.0,
        },
        "riskMetrics": {
            "annualizedVolatility": vol,
            "volatilityPercent": f"{vol * 100:.2f}%",
            "maxDrawdown": mdd["maxDrawdown"],
            "maxDrawdownPercent": f"{mdd['maxDrawdown'] * 100:.2f}%",
            "maxDrawdownPeakDate": mdd["peakDate"],
            "maxDrawdownTroughDate": mdd["troughDate"],
            "sharpeRatio": sharpe,
            "beta": beta,
        },
        "technicalIndicators": {
            "ma50": current_ma50,
            "ma200": current_ma200,
            "goldenCross": (current_ma50 is not None and current_ma200 is not None and current_ma50 > current_ma200),
            "priceVsMA50": (f"{((end_price - current_ma50) / current_ma50 * 100):.2f}%" if current_ma50 else None),
            "priceVsMA200": (f"{((end_price - current_ma200) / current_ma200 * 100):.2f}%" if current_ma200 else None),
        },
        "benchmarkComparison": {
            "correlation": corr,
            "correlationPercent": f"{corr * 100:.2f}%",
            "benchmarkCAGR": bench_cagr,
            "outperformance": outperformance,
        },
        "rawData": {
            "prices": prices[-1000:],   # prevent excessive payload
            "returns": rets[-1000:],
            "ma50": ma50[-1000:],
            "ma200": ma200[-1000:],
        },
    }


# -----------------------------
# Financial Ratio Extraction (heuristic)
# -----------------------------

_FIN_PATTERNS = {
    "revenue": re.compile(r"(?:pendapatan|revenue|penjualan|sales)[\s\S]{0,120}?([\d\.,]+)", re.IGNORECASE),
    "cogs": re.compile(r"(?:beban pokok|cost of goods|harga pokok)[\s\S]{0,120}?([\d\.,]+)", re.IGNORECASE),
    "opex": re.compile(r"(?:beban operasi|operating expense|beban usaha)[\s\S]{0,120}?([\d\.,]+)", re.IGNORECASE),
    "net_profit": re.compile(r"(?:laba bersih|net profit|net income|laba tahun berjalan)[\s\S]{0,120}?([\d\.,]+)", re.IGNORECASE),
    "total_assets": re.compile(r"(?:total aset|total assets|jumlah aset)[\s\S]{0,120}?([\d\.,]+)", re.IGNORECASE),
    "total_liabilities": re.compile(r"(?:total liabilitas|total liabilities|total kewajiban)[\s\S]{0,120}?([\d\.,]+)", re.IGNORECASE),
    "equity": re.compile(r"(?:total ekuitas|total equity|modal)[\s\S]{0,120}?([\d\.,]+)", re.IGNORECASE),
    "total_debt": re.compile(r"(?:total utang|total debt|pinjaman)[\s\S]{0,120}?([\d\.,]+)", re.IGNORECASE),
    "interest_expense": re.compile(r"(?:beban bunga|interest expense|biaya bunga)[\s\S]{0,120}?([\d\.,]+)", re.IGNORECASE),
}

def _extract_number(text: str, pat: re.Pattern) -> float:
    m = pat.search(text)
    if not m:
        return 0.0
    s = m.group(1)
    s = s.replace(" ", "").replace(",", "")
    # many Indonesian statements use dot as thousand separator; heuristic:
    # if there are multiple dots and no comma, remove dots
    if s.count(".") >= 2 and "," not in s:
        s = s.replace(".", "")
    return safe_float(s, 0.0)


def extract_financial_ratios_from_texts(texts: List[str], years: int = 5) -> Dict[str, Any]:
    """
    Produces a 5-year array of metrics with ratios + qualitative notes.
    Year labels are derived as (current year - (years-1-i)).
    """
    current_year = now_jakarta().year
    out: List[Dict[str, Any]] = []

    for i, text in enumerate(texts[:years]):
        year = current_year - (years - 1 - i)
        t = text or ""

        revenue = _extract_number(t, _FIN_PATTERNS["revenue"])
        cogs = _extract_number(t, _FIN_PATTERNS["cogs"])
        opex = _extract_number(t, _FIN_PATTERNS["opex"])
        net_profit = _extract_number(t, _FIN_PATTERNS["net_profit"])
        total_assets = _extract_number(t, _FIN_PATTERNS["total_assets"])
        total_liab = _extract_number(t, _FIN_PATTERNS["total_liabilities"])
        equity = _extract_number(t, _FIN_PATTERNS["equity"]) or max(total_assets - total_liab, 0.0)
        total_debt = _extract_number(t, _FIN_PATTERNS["total_debt"])
        interest = _extract_number(t, _FIN_PATTERNS["interest_expense"])

        gross_profit = max(revenue - cogs, 0.0)
        ebit = max(gross_profit - opex, 0.0)

        gross_margin = (gross_profit / revenue * 100) if revenue else 0.0
        net_margin = (net_profit / revenue * 100) if revenue else 0.0
        roa = (net_profit / total_assets * 100) if total_assets else 0.0
        roe = (net_profit / equity * 100) if equity else 0.0
        debt_to_equity = (total_debt / equity) if equity else 0.0
        interest_cov = (ebit / interest) if interest else 0.0

        # Qualitative notes (truncate)
        def note(pattern: str, default: str) -> str:
            m = re.search(pattern, t, flags=re.IGNORECASE | re.DOTALL)
            if not m:
                return default
            return (m.group(0)[:500]).strip()

        out.append({
            "year": year,
            "revenue": revenue,
            "cogs": cogs,
            "grossProfit": gross_profit,
            "operatingExpenses": opex,
            "ebit": ebit,
            "interestExpense": interest,
            "netProfit": net_profit,
            "totalAssets": total_assets,
            "totalLiabilities": total_liab,
            "totalDebt": total_debt,
            "equity": equity,
            "ratios": {
                "grossMarginPct": round(gross_margin, 2),
                "netMarginPct": round(net_margin, 2),
                "roaPct": round(roa, 2),
                "roePct": round(roe, 2),
                "debtToEquity": round(debt_to_equity, 2),
                "interestCoverage": round(interest_cov, 2),
            },
            "notes": {
                "debtMaturity": note(r"(jatuh tempo|maturity|pelunasan utang)[\s\S]{0,500}", "Not found"),
                "capex": note(r"(belanja modal|capital expenditure|capex|investasi aset)[\s\S]{0,500}", "Not found"),
                "litigation": note(r"(litigasi|litigation|gugatan|perkara hukum)[\s\S]{0,500}", "No litigation mentioned"),
                "esg": note(r"(keberlanjutan|sustainability|esg|lingkungan|sosial|tata kelola)[\s\S]{0,500}", "Not found"),
                "corporateActions": note(r"(aksi korporasi|corporate action|dividen|stock split|rights issue)[\s\S]{0,500}", "Not found"),
            }
        })

    trends = {}
    if len(out) >= 2:
        def growth(first: float, last: float) -> float:
            return ((last - first) / first * 100) if first else 0.0
        trends = {
            "revenueGrowthPct": growth(out[0]["revenue"], out[-1]["revenue"]),
            "netProfitGrowthPct": growth(out[0]["netProfit"], out[-1]["netProfit"]),
            "assetGrowthPct": growth(out[0]["totalAssets"], out[-1]["totalAssets"]),
            "avgROEPct": statistics.fmean([x["ratios"]["roePct"] for x in out]) if out else 0.0,
            "avgROAPct": statistics.fmean([x["ratios"]["roaPct"] for x in out]) if out else 0.0,
            "avgNetMarginPct": statistics.fmean([x["ratios"]["netMarginPct"] for x in out]) if out else 0.0,
        }

    return {"financialData": out, "trends": trends, "extractedAt": now_jakarta().isoformat()}


# -----------------------------
# News -> SWOT Classification (heuristic)
# -----------------------------

_SWOT_KEYWORDS = {
    "strength": ["pertumbuhan", "meningkat", "ekspansi", "inovasi", "penghargaan", "prestasi", "keunggulan", "unggul", "sukses", "profit", "laba", "dividen", "akuisisi", "kemitraan", "kolaborasi"],
    "weakness": ["penurunan", "turun", "merosot", "kerugian", "rugi", "masalah", "kendala", "hambatan", "defisit", "utang", "gagal", "ditunda", "dibatalkan"],
    "opportunity": ["peluang", "potensi", "prospek", "rencana", "target", "proyeksi", "akan", "investasi", "pasar baru", "ekspor", "digitalisasi", "transformasi"],
    "threat": ["risiko", "ancaman", "tantangan", "persaingan", "kompetitor", "regulasi", "sanksi", "denda", "investigasi", "gugatan", "krisis", "inflasi", "resesi"],
}
_THEME_KEYWORDS = {
    "regulasi": ["regulasi", "peraturan", "kebijakan", "ojk", "pemerintah", "undang-undang", "aturan", "izin", "compliance"],
    "kompetisi": ["kompetitor", "pesaing", "persaingan", "market share", "pangsa pasar", "rival"],
    "teknologi": ["teknologi", "digital", "inovasi", "aplikasi", "platform", "sistem", "otomasi", "ai", "blockchain"],
    "operasional": ["operasional", "produksi", "kapasitas", "efisiensi", "supply chain", "distribusi", "logistik"],
    "keuangan": ["keuangan", "laba", "rugi", "revenue", "pendapatan", "profit", "dividen", "utang", "modal", "investasi"],
    "manajemen": ["manajemen", "direksi", "komisaris", "ceo", "cfo", "kepemimpinan", "strategi", "restrukturisasi"],
    "reputasi": ["reputasi", "citra", "brand", "merek", "penghargaan", "skandal", "kontroversi"],
    "esg": ["esg", "lingkungan", "sosial", "governance", "sustainability", "berkelanjutan", "emisi", "csr"],
    "makro": ["ekonomi", "inflasi", "suku bunga", "nilai tukar", "gdp", "resesi", "pertumbuhan ekonomi", "makro"],
}

_POS_WORDS = ["baik", "positif", "meningkat", "tumbuh", "sukses", "untung", "laba", "naik", "optimis", "kuat"]
_NEG_WORDS = ["buruk", "negatif", "turun", "merosot", "rugi", "gagal", "lemah", "pesimis", "krisis", "masalah"]
_HIGH_IMPACT = ["signifikan", "besar", "utama", "penting", "krusial", "strategis", "fundamental", "drastis"]
_FIN_WORDS = ["miliar", "triliun", "juta", "persen", "%", "revenue", "laba", "rugi"]


def _count_matches(text: str, keywords: List[str]) -> int:
    t = text.lower()
    return sum(1 for kw in keywords if kw.lower() in t)


def classify_news_to_swot(news: List[NewsItem]) -> Dict[str, Any]:
    classified: List[ClassifiedNewsItem] = []

    for n in news:
        title = n.title or ""
        desc = n.snippet or ""
        text = f"{title} {desc}".lower()

        scores = {k: _count_matches(text, v) for k, v in _SWOT_KEYWORDS.items()}
        swot = max(scores, key=scores.get) if max(scores.values()) > 0 else "opportunity"

        theme_scores = {k: _count_matches(text, v) for k, v in _THEME_KEYWORDS.items()}
        theme = max(theme_scores, key=theme_scores.get) if max(theme_scores.values()) > 0 else "operasional"

        pos = _count_matches(text, _POS_WORDS)
        neg = _count_matches(text, _NEG_WORDS)
        sentiment = (pos - neg) / max(pos + neg, 1)
        if swot == "strength":
            sentiment = max(sentiment, 0.3)
        if swot in {"weakness", "threat"}:
            sentiment = min(sentiment, -0.2)
        if swot == "opportunity":
            sentiment = max(sentiment, 0.1)
        sentiment = max(-1.0, min(1.0, sentiment))

        score = 5.0
        score += _count_matches(text, _HIGH_IMPACT) * 1.5
        score += _count_matches(text, _FIN_WORDS) * 1.0
        theme_w = {"keuangan": 1.5, "regulasi": 1.3, "manajemen": 1.2, "operasional": 1.0, "kompetisi": 1.1, "teknologi": 0.9, "reputasi": 1.0, "esg": 0.8, "makro": 1.2}
        score *= theme_w.get(theme, 1.0)
        score = max(0.0, min(10.0, round(score, 1)))

        classified.append(ClassifiedNewsItem(
            title=title,
            description=desc,
            link=n.url,
            pubDate=n.date.isoformat(),
            swotCategory=swot,
            theme=theme,
            sentimentScore=sentiment,
            materialityScore=score,
        ))

    classified.sort(key=lambda x: x.materialityScore, reverse=True)
    return {"classifiedNews": [c.model_dump() for c in classified], "totalNews": len(classified)}


# -----------------------------
# Bond Data (Optional)
# -----------------------------

def fetch_bond_data_optional(isins: List[str]) -> Dict[str, Any]:
    """
    Mirrors the n8n placeholder.
    Provide BOND_API_URL and any required auth headers/token via env if you have a public bond API.
    """
    if not isins:
        return {"status": "skipped", "reason": "no ISINs provided"}

    url = os.getenv("BOND_API_URL", "").strip()
    if not url or "PLACEHOLDER" in url:
        return {"status": "skipped", "reason": "BOND_API_URL not configured"}

    params = {"isin": ",".join(isins)}
    headers = {}
    token = os.getenv("BOND_API_TOKEN")
    if token:
        headers["Authorization"] = f"Bearer {token}"

    r = requests.get(url, params=params, headers=headers, timeout=60)
    r.raise_for_status()
    return {"status": "ok", "data": r.json()}


# -----------------------------
# OpenAI Big-4 report (Structured Outputs)
# -----------------------------

def generate_big4_report_openai(
    config: AnalysisRequest,
    stock_analytics: Dict[str, Any],
    financial_ratios: Dict[str, Any],
    swot_news: Dict[str, Any],
    bond_data: Dict[str, Any],
    model: str = "gpt-4o-mini",
) -> Big4Report:
    """
    Uses OpenAI Chat Completions with Structured Outputs (JSON Schema).
    Requires: OPENAI_API_KEY

    If you prefer the Responses API, you can adapt this function accordingly.
    """
    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY not set")

    from openai import OpenAI  # openai>=1.x

    system_message = (
        "Anda adalah analis keuangan senior dari firma konsultan Big-4. "
        "Tugas Anda: buat laporan analisis emiten Indonesia (konteks OJK/IDX) "
        "berbasis data 5 tahun terakhir dalam Bahasa Indonesia.\n\n"
        "Gaya penulisan Big-4:\n"
        "- Heading jelas dengan struktur hierarki\n"
        "- Bullet points untuk key messages\n"
        "- Evidence-based: setiap poin SWOT harus disertai bukti dari data\n"
        "- Sebutkan asumsi dan keterbatasan data secara eksplisit\n"
        "- Profesional, objektif, tidak memberikan rekomendasi investasi eksplisit "
        "(tidak ada buy/sell/target price)\n\n"
        "Struktur laporan harus sesuai skema JSON yang diberikan."
    )

    user_message = (
        f"Analisa data berikut untuk perusahaan {config.company_name} ({config.stock_ticker}):\n"
        f"Data Saham 5Y: {json.dumps(stock_analytics, ensure_ascii=False)}\n"
        f"Data Keuangan 5Y: {json.dumps(financial_ratios, ensure_ascii=False)}\n"
        f"Analisis SWOT dari Berita: {json.dumps(swot_news, ensure_ascii=False)}\n"
        f"Data Obligasi: {json.dumps(bond_data, ensure_ascii=False)}\n\n"
        "Keluaran harus valid JSON dan sesuai skema."
    )

    schema = Big4Report.model_json_schema()

    client = OpenAI(api_key=api_key)

    resp = client.chat.completions.create(
        model=model,
        temperature=0.3,
        messages=[
            {"role": "system", "content": system_message},
            {"role": "user", "content": user_message},
        ],
        response_format={
            "type": "json_schema",
            "json_schema": {
                "name": "big4_report",
                "strict": True,
                "schema": schema,
            },
        },
    )

    content = resp.choices[0].message.content or "{}"
    try:
        data = json.loads(content)
    except json.JSONDecodeError as e:
        raise RuntimeError(f"OpenAI output was not valid JSON: {e}\nContent: {content[:500]}") from e

    try:
        return Big4Report.model_validate(data)
    except ValidationError as e:
        raise RuntimeError(f"OpenAI JSON did not match schema: {e}") from e


# -----------------------------
# Report Formatting: PDF + DOCX
# -----------------------------

def render_pdf(report: Big4Report, out_path: Path, title: str) -> None:
    """
    Uses reportlab (already available in this environment, but install in your runtime if needed).
    """
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle

    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=styles["Heading1"], textColor=colors.HexColor("#003366"))
    h2 = ParagraphStyle("h2", parent=styles["Heading2"], textColor=colors.HexColor("#0066CC"))
    body = styles["BodyText"]

    doc = SimpleDocTemplate(str(out_path), pagesize=A4, rightMargin=2*cm, leftMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    story: List[Any] = []

    story.append(Paragraph(title, h1))
    story.append(Paragraph(f"Tanggal: {now_jakarta().strftime('%d %B %Y')}", body))
    story.append(Spacer(1, 12))

    def add_section(name: str, content: str):
        story.append(Paragraph(name, h2))
        for para in (content or "").split("\n"):
            if para.strip():
                story.append(Paragraph(para.strip(), body))
        story.append(Spacer(1, 10))

    add_section("1. Executive Summary", report.executive_summary)
    add_section("2. Company Overview", report.company_overview)
    add_section("3. Industry", report.industry)
    add_section("4. Stock Performance", report.stock_performance.narrative)
    story.append(Paragraph("Key Metrics (Stock)", h2))
    story.append(Paragraph(json.dumps(report.stock_performance.metrics, ensure_ascii=False, indent=2), styles["Code"]))
    story.append(Spacer(1, 10))

    add_section("5. Bond View", report.bond_view.narrative + "\n\nLimitations:\n" + report.bond_view.limitations)
    add_section("6. Financial Analysis", report.financial_analysis.highlights)

    story.append(Paragraph("7. SWOT", h2))
    for label, arr in [("Strengths", report.swot.S), ("Weaknesses", report.swot.W), ("Opportunities", report.swot.O), ("Threats", report.swot.T)]:
        story.append(Paragraph(label, styles["Heading3"]))
        bullets = ListFlowable([ListItem(Paragraph(x, body)) for x in (arr or ["(Tidak ada data)"])], bulletType="bullet")
        story.append(bullets)
        story.append(Spacer(1, 8))

    story.append(Paragraph("8. Key Risks", h2))
    bullets = ListFlowable([ListItem(Paragraph(x, body)) for x in (report.key_risks or ["(Tidak ada data)"])], bulletType="bullet")
    story.append(bullets)
    story.append(Spacer(1, 10))

    add_section("9. Governance, ESG & Regulatory Notes", report.governance_esg_regulatory_notes)

    add_section("10. Appendix - Methodology", report.appendix.methodology)
    story.append(Paragraph("Appendix - Sources", styles["Heading3"]))
    bullets = ListFlowable([ListItem(Paragraph(x, body)) for x in (report.appendix.sources or ["(Tidak ada data)"])], bulletType="bullet")
    story.append(bullets)
    story.append(Spacer(1, 6))
    story.append(Paragraph("Appendix - Assumptions", styles["Heading3"]))
    bullets = ListFlowable([ListItem(Paragraph(x, body)) for x in (report.appendix.assumptions or ["(Tidak ada data)"])], bulletType="bullet")
    story.append(bullets)

    doc.build(story)


def render_docx(report: Big4Report, out_path: Path, title: str) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(f"Tanggal: {now_jakarta().strftime('%d %B %Y')}")

    def add_heading(text: str, level: int = 2):
        doc.add_heading(text, level=level)

    def add_paras(text: str):
        for para in (text or "").split("\n"):
            if para.strip():
                doc.add_paragraph(para.strip())

    add_heading("1. Executive Summary"); add_paras(report.executive_summary)
    add_heading("2. Company Overview"); add_paras(report.company_overview)
    add_heading("3. Industry"); add_paras(report.industry)
    add_heading("4. Stock Performance"); add_paras(report.stock_performance.narrative)

    add_heading("Key Metrics (Stock)", level=3)
    doc.add_paragraph(json.dumps(report.stock_performance.metrics, ensure_ascii=False, indent=2))

    add_heading("5. Bond View"); add_paras(report.bond_view.narrative)
    add_heading("Limitations", level=3); add_paras(report.bond_view.limitations)

    add_heading("6. Financial Analysis"); add_paras(report.financial_analysis.highlights)

    add_heading("7. SWOT")
    for label, arr in [("Strengths", report.swot.S), ("Weaknesses", report.swot.W), ("Opportunities", report.swot.O), ("Threats", report.swot.T)]:
        add_heading(label, level=3)
        for x in (arr or ["(Tidak ada data)"]):
            doc.add_paragraph(x, style="List Bullet")

    add_heading("8. Key Risks")
    for x in (report.key_risks or ["(Tidak ada data)"]):
        doc.add_paragraph(x, style="List Bullet")

    add_heading("9. Governance, ESG & Regulatory Notes"); add_paras(report.governance_esg_regulatory_notes)

    add_heading("10. Appendix", level=2)
    add_heading("Methodology", level=3); add_paras(report.appendix.methodology)
    add_heading("Sources", level=3)
    for x in (report.appendix.sources or ["(Tidak ada data)"]):
        doc.add_paragraph(x, style="List Bullet")
    add_heading("Assumptions", level=3)
    for x in (report.appendix.assumptions or ["(Tidak ada data)"]):
        doc.add_paragraph(x, style="List Bullet")

    doc.save(str(out_path))


# -----------------------------
# Google Drive + Sheets logging (Optional)
# -----------------------------

def upload_to_drive(file_path: Path, folder_id: str) -> Dict[str, Any]:
    """
    Uploads to Google Drive and returns {fileId, webViewLink}.
    Requires Google Drive API permissions.
    """
    creds = _get_google_creds()
    from googleapiclient.discovery import build
    from googleapiclient.http import MediaFileUpload

    drive = build("drive", "v3", credentials=creds, cache_discovery=False)
    media = MediaFileUpload(str(file_path), resumable=True)
    body = {"name": file_path.name, "parents": [folder_id]}

    created = drive.files().create(
        body=body,
        media_body=media,
        fields="id, webViewLink, webContentLink",
    ).execute()
    return created


def log_execution_to_sheets(sheet_id: str, sheet_name: str, row: List[Any]) -> None:
    creds = _get_google_creds()
    from googleapiclient.discovery import build
    service = build("sheets", "v4", credentials=creds, cache_discovery=False)

    rng = f"{sheet_name}!A:Z"
    body = {"values": [row]}
    service.spreadsheets().values().append(
        spreadsheetId=sheet_id,
        range=rng,
        valueInputOption="USER_ENTERED",
        insertDataOption="INSERT_ROWS",
        body=body,
    ).execute()


# -----------------------------
# Orchestration
# -----------------------------

def run_pipeline(req: AnalysisRequest) -> Dict[str, Any]:
    req = req.normalized()
    rid = run_id(req.stock_ticker)

    with tempfile.TemporaryDirectory(prefix="analisa-emiten-") as td:
        workdir = Path(td)
        outdir = Path.cwd() / "outputs"
        outdir.mkdir(parents=True, exist_ok=True)

        # 1) Stock data
        stock_rows = []
        try:
            stock_rows = read_stock_data_google_sheets(req.google_sheet_id, "Stock Data")
        except Exception as e:
            LOGGER.warning("Failed reading Google Sheets stock data: %s", e)

        stock_points, stock_stats = parse_and_clean_stock_data(stock_rows)

        # Benchmark data: optional (if you store it in another sheet called 'Benchmark Data')
        benchmark_points: Optional[List[StockPoint]] = None
        try:
            bench_rows = read_stock_data_google_sheets(req.google_sheet_id, "Benchmark Data")
            benchmark_points, _ = parse_and_clean_stock_data(bench_rows)
        except Exception:
            benchmark_points = None

        stock_analytics = calculate_stock_analytics(stock_points, req.period_years, benchmark_points)

        # 2) News
        news_items = fetch_google_news_rss(req.news_query or req.company_name, limit=80)
        swot_news = classify_news_to_swot(news_items)

        # 3) Annual reports
        annual_texts: List[str] = []
        annual_paths: List[Path] = []
        annual_limitations = None
        if req.annual_report_urls:
            annual_paths = download_files(req.annual_report_urls, workdir / "annual_reports", max_files=5)
            annual_texts = extract_text_from_pdfs(annual_paths)
        else:
            annual_limitations = "annual_report_urls kosong; analisis rasio keuangan berbasis annual report dilewati."

        financial_ratios = extract_financial_ratios_from_texts(annual_texts, years=min(5, len(annual_texts))) if annual_texts else {
            "financialData": [],
            "trends": {},
            "extractedAt": now_jakarta().isoformat(),
            "limitations": annual_limitations,
        }

        # 4) Bond data optional
        bond_data = fetch_bond_data_optional(req.bond_identifiers)

        # 5) Big-4 report (OpenAI)
        report = generate_big4_report_openai(
            req,
            stock_analytics=stock_analytics,
            financial_ratios=financial_ratios,
            swot_news=swot_news,
            bond_data=bond_data,
        )

        # 6) Persist artifacts
        report_json_path = outdir / f"{rid}_report.json"
        report_pdf_path = outdir / f"{rid}_Big4Report.pdf"
        report_docx_path = outdir / f"{rid}_Big4Report.docx"
        raw_data_path = outdir / f"{rid}_RawData.json"

        report_json_path.write_text(report.model_dump_json(indent=2, ensure_ascii=False), encoding="utf-8")

        title = f"Analisis Komprehensif Emiten - {req.company_name} ({req.stock_ticker})"
        render_pdf(report, report_pdf_path, title=title)
        render_docx(report, report_docx_path, title=title)

        raw_data = {
            "run_id": rid,
            "config": req.model_dump(),
            "stock_stats": stock_stats,
            "stock_analytics": stock_analytics,
            "financial_ratios": financial_ratios,
            "swot_news": swot_news,
            "bond_data": bond_data,
            "news_count": len(news_items),
            "annual_report_files": [p.name for p in annual_paths],
            "generated_at": now_jakarta().isoformat(),
        }
        raw_data_path.write_text(json.dumps(raw_data, ensure_ascii=False, indent=2), encoding="utf-8")

        # 7) Optional: upload to Drive
        drive_links: Dict[str, Any] = {"report_pdf": None, "raw_data": None}
        if req.drive_folder_id and "PLACEHOLDER" not in req.drive_folder_id:
            try:
                uploaded_pdf = upload_to_drive(report_pdf_path, req.drive_folder_id)
                uploaded_raw = upload_to_drive(raw_data_path, req.drive_folder_id)
                drive_links = {
                    "report_pdf": uploaded_pdf.get("webViewLink"),
                    "raw_data": uploaded_raw.get("webViewLink"),
                }
            except Exception as e:
                LOGGER.warning("Drive upload failed: %s", e)

        # 8) Optional: log to Sheets
        if req.google_sheet_id and "PLACEHOLDER" not in req.google_sheet_id:
            try:
                log_execution_to_sheets(
                    req.google_sheet_id,
                    "Execution Log",
                    [
                        rid,
                        now_jakarta().isoformat(),
                        req.company_name,
                        req.stock_ticker,
                        "success",
                        len(stock_points),
                        len(news_items),
                        "PDF/DOCX/JSON",
                        drive_links.get("report_pdf") or "",
                    ],
                )
            except Exception as e:
                LOGGER.warning("Sheets log failed: %s", e)

        return {
            "status": "success",
            "run_id": rid,
            "outputs": {
                "report_pdf_path": str(report_pdf_path),
                "report_docx_path": str(report_docx_path),
                "report_json_path": str(report_json_path),
                "raw_data_path": str(raw_data_path),
            },
            "drive_links": drive_links,
            "key_metrics": {
                "cagr_5y": stock_analytics.get("returnMetrics", {}).get("cagr"),
                "volatility": stock_analytics.get("riskMetrics", {}).get("annualizedVolatility"),
                "avg_roe": (financial_ratios.get("trends", {}).get("avgROEPct") if isinstance(financial_ratios, dict) else None),
            },
            "limitations": {
                "annual_reports": annual_limitations,
                "bond_data": (bond_data.get("reason") if isinstance(bond_data, dict) else None),
            },
        }


# -----------------------------
# API server (Webhook + Schedule)
# -----------------------------

def build_app() -> Any:
    if FastAPI is None:
        raise RuntimeError("FastAPI/Apscheduler not installed. Install requirements and retry.")

    app = FastAPI(title="Analisis Data Perusahaan Tbk", version="1.0.0")

    @app.post("/analisa-emiten")
    def analisa_emiten(payload: Dict[str, Any]):
        try:
            req = AnalysisRequest.model_validate(payload)
        except ValidationError as e:
            raise HTTPException(status_code=400, detail=json.loads(e.json()))
        try:
            result = run_pipeline(req)
            return JSONResponse(result)
        except Exception as e:
            LOGGER.exception("Pipeline failed")
            raise HTTPException(status_code=500, detail=str(e))

    # Weekly schedule: Monday 06:00 Asia/Jakarta
    scheduler = BackgroundScheduler(timezone="Asia/Jakarta")

    @app.on_event("startup")
    def _startup():
        # If you want a weekly job, define defaults via env var JSON
        default_payload = os.getenv("DEFAULT_ANALYSIS_PAYLOAD_JSON", "").strip()
        if default_payload:
            try:
                default_req = AnalysisRequest.model_validate(json.loads(default_payload))
            except Exception as e:
                LOGGER.warning("Invalid DEFAULT_ANALYSIS_PAYLOAD_JSON; schedule disabled: %s", e)
                return

            trigger = CronTrigger(day_of_week="mon", hour=6, minute=0)
            scheduler.add_job(lambda: run_pipeline(default_req), trigger=trigger, id="weekly-analysis", replace_existing=True)
            scheduler.start()
            LOGGER.info("Weekly scheduler started (Mon 06:00 Asia/Jakarta).")

    @app.on_event("shutdown")
    def _shutdown():
        try:
            scheduler.shutdown(wait=False)
        except Exception:
            pass

    return app


# Expose for uvicorn: `uvicorn analisis_data_perusahaan_tbk_py311:app --reload`
app = build_app() if FastAPI is not None else None


# -----------------------------
# CLI
# -----------------------------

def main():
    logging.basicConfig(level=os.getenv("LOG_LEVEL", "INFO"))
    parser = argparse.ArgumentParser(description="Analisis Data Perusahaan Tbk (Python 3.11)")
    parser.add_argument("--input-json", type=str, help="Path to a JSON payload (AnalysisRequest) to run once.")
    parser.add_argument("--serve", action="store_true", help="Run API server with webhook and weekly scheduler.")
    parser.add_argument("--host", type=str, default="0.0.0.0")
    parser.add_argument("--port", type=int, default=8000)
    args = parser.parse_args()

    if args.serve:
        if FastAPI is None:
            raise RuntimeError("fastapi/uvicorn/apscheduler not installed")
        import uvicorn
        uvicorn.run("analisis_data_perusahaan_tbk_py311:app", host=args.host, port=args.port, reload=False)
        return

    if not args.input_json:
        raise SystemExit("Provide --input-json to run once, or --serve to start the webhook server.")

    payload = json.loads(Path(args.input_json).read_text(encoding="utf-8"))
    req = AnalysisRequest.model_validate(payload)
    result = run_pipeline(req)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
